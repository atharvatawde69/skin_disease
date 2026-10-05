"""Build the mini-project presentation from the college PPT template.

Usage:  python report/build_ppt.py
Needs:  report/ppt_template/template.pptx  (made from the .ppt template by report/ppt_inspect.ps1),
        report/figures/*.png               (made by make_figures.py and ppt_figures.py)
Writes: report/Skin_Disease_Detection_Presentation.pptx
Yellow-highlighted text marks placeholders to fill in by hand (SDG, group number, members).
"""
import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
FIG = REPORT / "figures"
TEMPLATE = REPORT / "ppt_template" / "template.pptx"
OUT = REPORT / "Skin_Disease_Detection_Presentation.pptx"

FONT = "Times New Roman"          # the template's font
INK, DARK, ACCENT, TINT, RED, RED_TINT = "1A1A1A", "1F5E8C", "21A7DD", "E4F1F9", "B03A2E", "F8E9E6"
GREY = "5A6770"
X0, WIDTH = 0.5, 7.4              # the blue artwork covers the right part of the slide: stay inside this strip

prs = Presentation(str(TEMPLATE))
LAYOUT = prs.slide_layouts.get_by_name("Title and Content")
BG = copy.deepcopy(prs.slides[3]._element.find(qn("p:cSld")).find(qn("p:bg")))
old_slides = list(prs.slides)


# ------------------------------------------------------------------ helpers
def rgb(hex_):
    return RGBColor.from_string(hex_)


def add_highlight(run):
    rpr = run._r.get_or_add_rPr()
    hl = etree.SubElement(rpr, qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr")).set("val", "FFFF00")
    rpr.remove(hl)
    anchor = next((c for c in rpr if c.tag in (qn("a:latin"), qn("a:ea"), qn("a:cs"), qn("a:sym"), qn("a:hlinkClick"))), None)
    if anchor is not None:
        anchor.addprevious(hl)
    else:
        rpr.append(hl)


def set_bullet(par, indent_in=0.3):
    ppr = par._p.get_or_add_pPr()
    ppr.set("marL", str(int(Inches(indent_in))))
    ppr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    bu_font = etree.SubElement(ppr, qn("a:buFont"))
    bu_font.set("typeface", "Arial")
    etree.SubElement(ppr, qn("a:buChar")).set("char", "•")


def write(par, runs, size, bold=False, color=INK, italic=False):
    """runs: a string, or a list of strings / (text, {bold, color, hl, italic}) tuples."""
    if isinstance(runs, str):
        runs = [runs]
    for item in runs:
        text, opt = (item, {}) if isinstance(item, str) else item
        r = par.add_run()
        r.text = text
        f = r.font
        f.name = FONT
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.italic = opt.get("italic", italic)
        f.color.rgb = rgb(opt.get("color", color))
        if opt.get("hl"):
            add_highlight(r)


def textbox(slide, x, y, w, h, name, anchor=MSO_ANCHOR.TOP, margin=0.05):
    shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shp.name = name
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    return shp, tf


def paragraphs(tf, items, size, bullet=False, after=6, align=PP_ALIGN.LEFT, color=INK, bold=False):
    """items: list of run-lists. Replaces the (empty) first paragraph, then adds the rest."""
    for i, runs in enumerate(items):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_after = Pt(after)
        par.space_before = Pt(0)
        if align is not None:
            par.alignment = align
        write(par, runs, size, bold=bold, color=color)
        if bullet:
            set_bullet(par)


def card(slide, x, y, w, h, name, fill=TINT, line=None, radius=0.05):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.name = name
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.1)
    return shp, tf


def new_slide(title, notes):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    s._element.find(qn("p:cSld")).insert(0, copy.deepcopy(BG))
    _, tf = textbox(s, X0, 0.3, WIDTH, 1.2, "Title", anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
    paragraphs(tf, [[title]], 36, after=0, bold=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def picture(slide, filename, x, y, w, name, alt):
    pic = slide.shapes.add_picture(str(FIG / filename), Inches(x), Inches(y), width=Inches(w))
    pic.name = name
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def fill_cell(cell, runs, size, bold=False, color=INK, fill=None, align=PP_ALIGN.LEFT):
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.05)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    par = tf.paragraphs[0]
    par.alignment = align
    write(par, runs, size, bold=bold, color=color)
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)


def make_table(slide, x, y, widths, header, rows, size, name, row_h=0.6):
    shp = slide.shapes.add_table(len(rows) + 1, len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * (len(rows) + 1)))
    shp.name = name
    tbl = shp.table
    tbl.horz_banding = False
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for r in tbl.rows:
        r.height = Inches(row_h)
    for j, h in enumerate(header):
        fill_cell(tbl.cell(0, j), h, size, bold=True, color="FFFFFF", fill=DARK)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            fill_cell(tbl.cell(i, j), val, size, bold=(j == 0), fill=TINT if i % 2 == 0 else "FFFFFF")
    return shp


def style_chart(chart, size=17):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(INK)


# ================================================================== slide 1: title (template slide 2)
title_slide = old_slides[1]
for shp in list(title_slide.shapes):
    if shp.name in ("Title 7", "TextBox 1"):
        shp._element.getparent().remove(shp._element)
_, tf = textbox(title_slide, 0.55, 2.0, 9.2, 1.1, "Project title", anchor=MSO_ANCHOR.MIDDLE)
paragraphs(tf, [["Skin Disease Detection using Deep Learning"]], 32, after=0, align=PP_ALIGN.CENTER, bold=True)
_, tf = textbox(title_slide, 0.9, 3.3, 8.6, 3.6, "Project details")
HL = {"hl": True}
paragraphs(tf, [
    [("Theme- ", {"bold": True}), ("SDG [add number and name]", HL)],
    [("Category- ", {"bold": True}), "Software"],
    [("Group No- ", {"bold": True}), ("[add]", HL)],
    [("Group Members- ", {"bold": True}), ("[add names]", HL)],
    [("Subject Name- ", {"bold": True}), "Deep Learning Lab (CSL701)"],
    [("Subject Incharge- ", {"bold": True}), "Prof. Vijaya Bharathi J"],
], 22, bullet=True, after=8)
title_slide.notes_slide.notes_text_frame.text = (
    "Good morning. Our project is skin disease detection using deep learning. We built a system that looks at a "
    "close-up image of a skin lesion, tells which of seven types it most likely is, and shows how sure it is.")

# ================================================================== slide 2: introduction
s = new_slide("Introduction",
              "Skin cancer is one of the most common cancers, and finding it early makes a big difference. But "
              "dermatologists are not available everywhere. A quick second opinion from software can help people decide "
              "when to see a doctor. Our objective is to classify a lesion image into seven types. "
              "[Mention here the SDG this project supports.]")
_, tf = textbox(s, X0, 1.65, WIDTH, 3.9, "Introduction points")
paragraphs(tf, [
    ["Skin cancer is one of the most common cancers"],
    ["Early detection greatly improves outcomes"],
    ["Specialists are not available everywhere"],
    [("Motivation: ", {"bold": True}), "a quick second opinion from software"],
    [("Objective: ", {"bold": True}), "classify a skin lesion image into 7 types"],
    [("SDG: ", {"bold": True}), ("[add the SDG mapped to this project]", HL)],
], 22, bullet=True, after=9)
tiles = [("10,015", "images"), ("7", "lesion types"), ("1", "web app")]
for i, (big, small) in enumerate(tiles):
    shp, tf = card(s, X0 + i * 2.5, 5.75, 2.4, 1.25, f"Stat tile {i + 1}")
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraphs(tf, [[big]], 30, after=0, align=PP_ALIGN.CENTER, bold=True, color=DARK)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    write(p, small, 18)

# ================================================================== slide 3: literature survey
s = new_slide("Literature Survey",
              "We reviewed five key works. Esteva and colleagues showed that a deep network can match dermatologists, but "
              "on two-class tasks. Tschandl's team released the HAM10000 dataset that we use. The ISIC 2018 challenge gave "
              "a benchmark. He's ResNet is the network we build on, and Grad-CAM is how we show what the model looks at.")
make_table(s, 0.35, 1.6, [1.85, 2.15, 2.0, 2.3, 2.0],
           ["Author & year", "Technique", "Dataset", "Key result", "Limitation"],
           [["Esteva et al., 2017 [1]", "Inception v3, transfer learning", "129,450 clinical images", "Matched 21 dermatologists", "Two-class tasks only"],
            ["Tschandl et al., 2018 [2]", "Dataset paper (no model)", "HAM10000: 10,015 images", "Over 50% biopsy-confirmed", "Many images per lesion"],
            ["Codella et al., 2019 [3]", "ISIC 2018 challenge", "12,500+ images", "159 teams in classification", "Equal scores generalise differently"],
            ["He et al., 2016 [4]", "ResNet, shortcut connections", "ImageNet", "3.57% error; won ILSVRC 2015", "Built for natural images"],
            ["Selvaraju et al., 2017 [5]", "Grad-CAM heatmaps", "ILSVRC-15 and others", "Beat earlier localisation methods", "Coarse maps"]],
           17, "Literature table", row_h=0.82)

# ================================================================== slide 4: limitations of current work
s = new_slide("Limitations of Current Work",
              "From these papers we found four gaps. Esteva's study used two-class tasks, so we classify all seven "
              "classes. The dataset has several photos of the same lesion, which can leak into the test set, so we split "
              "by lesion. Equal test scores do not mean equal quality, so we report balanced metrics and melanoma recall. "
              "And predictions usually come without an explanation, so we add heatmaps and warnings.")
make_table(s, X0, 1.75, [3.8, 3.6],
           ["Limitation found", "Our response"],
           [["Two-class tasks [1]", "Classify all 7 classes"],
            ["Several images of one lesion [2]", "Split the data by lesion"],
            ["Equal scores, different generalisation [3]", "Balanced metrics and melanoma recall"],
            ["Predictions without explanation [5]", "Heatmap and uncertainty warnings"]],
           19, "Limitations table", row_h=0.95)

# ================================================================== slide 5: problem statement and objectives
s = new_slide("Problem & Objectives",
              "The problem: given a dermoscopy image, predict one of seven classes, fairly, even though one class is two "
              "thirds of the data, and explain the result. Our objectives are to build a leak-free pipeline, compare a "
              "simple CNN with ResNet-18, test optimizers, regularization and loss functions, evaluate with balanced "
              "metrics, and deploy a web app.")
shp, tf = card(s, X0, 1.65, WIDTH, 1.95, "Problem statement")
paragraphs(tf, [[("Problem: ", {"bold": True, "color": DARK}),
                 "given a dermoscopy image of a skin lesion, predict which of 7 classes it belongs to, even though one "
                 "class is two thirds of the data, and explain the prediction."]], 20, after=0)
_, tf = textbox(s, X0, 3.8, WIDTH, 0.5, "Objectives heading")
paragraphs(tf, [["Objectives"]], 22, after=0, bold=True, color=DARK)
_, tf = textbox(s, X0, 4.3, WIDTH, 2.8, "Objectives list")
paragraphs(tf, [
    ["Build a leak-free pipeline (split by lesion)"],
    ["Compare a simple CNN with ResNet-18"],
    ["Test optimizers, regularization and loss"],
    ["Evaluate with balanced metrics and melanoma recall"],
    ["Deploy a web app with heatmap and guidance"],
], 20, bullet=True, after=6)

# ================================================================== slide 6: methodology, system diagram
s = new_slide("Proposed Methodology",
              "This is the whole system. On the left is training, done once on a Kaggle GPU: load the images, split by "
              "lesion, train and compare models, and keep the best one on the validation set. On the right is the web "
              "app: the user uploads an image, the saved ResNet-18 predicts seven probabilities, Grad-CAM draws the "
              "heatmap, and the app shows the result with guidance and warnings.")
picture(s, "ppt_system_diagram.png", X0, 1.75, WIDTH, "System diagram",
        "System diagram: training pipeline on the left, web application on the right, connected by the saved best model")

# ================================================================== slide 7: methodology, key ideas
s = new_slide("Methodology: Key Ideas",
              "Four ideas make this work. Split by lesion: all photos of one lesion stay in the same set, so test images "
              "are never seen in training. Transfer learning: we start from ResNet-18 already trained on ImageNet and then "
              "train it on skin images. Each block adds its input to its output, y equals F of x plus x, which makes deep "
              "networks easier to train. Class weights: rare classes count more in the loss, with weight N over K times n, "
              "where N is the number of images, K the number of classes and n the images in that class, so the model "
              "cannot just predict the common class. Grad-CAM: a heatmap of the regions that influenced the prediction. "
              "These design choices are our own contribution.")
ideas = [
    ("Split by lesion", "All photos of one lesion stay in one set. Test images are never seen in training."),
    ("Transfer learning", "Start from ResNet-18 trained on ImageNet [6], then train on skin images. Each block: y = F(x) + x"),
    ("Class weights", "Rare classes get a higher weight in the loss: w = N / (K × n)"),
    ("Grad-CAM heatmap", "Shows which parts of the image most influenced the prediction [5]"),
]
for i, (head, body_) in enumerate(ideas):
    x = X0 + (i % 2) * 3.8
    y = 1.65 + (i // 2) * 2.3
    shp, tf = card(s, x, y, 3.6, 2.15, f"Idea {i + 1}")
    paragraphs(tf, [[head]], 20, after=4, bold=True, color=DARK)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    write(p, body_, 17)
shp, tf = card(s, X0, 6.3, WIDTH, 0.75, "Our contribution", fill=RED_TINT)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
paragraphs(tf, [[("Our contribution: ", {"bold": True, "color": RED}),
                 "lesion-wise split, class-weighted training, uncertainty-aware app"]], 18, after=0)

# ================================================================== slide 8: experimental setup
s = new_slide("Experimental Setup",
              "We used HAM10000: 10,015 images in 7 classes, split by lesion into 7,054 training, 1,464 validation and "
              "1,497 test images. The bar chart shows the imbalance: common moles are two thirds of the data. We ran "
              "eleven experiments on a Tesla T4 GPU with PyTorch. Because of the imbalance we judge models by balanced "
              "accuracy, macro-F1 and melanoma recall, not by accuracy alone.")
_, tf = textbox(s, X0, 1.65, 3.85, 4.0, "Setup points")
paragraphs(tf, [
    [("Data: ", {"bold": True}), "HAM10000 [2], 10,015 images, 7 classes"],
    [("Split: ", {"bold": True}), "7,054 train, 1,464 val, 1,497 test"],
    [("Hardware: ", {"bold": True}), "Tesla T4 GPU (Kaggle), PyTorch"],
    [("Runs: ", {"bold": True}), "11 experiments"],
    [("Metrics: ", {"bold": True}), "balanced accuracy, macro-F1, melanoma recall"],
], 18, bullet=True, after=7)
cd = CategoryChartData()
cd.categories = ["nv", "mel", "bkl", "bcc", "akiec", "vasc", "df"]
cd.add_series("Images", (6705, 1113, 1099, 514, 327, 142, 115))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(4.45), Inches(1.6), Inches(3.5), Inches(4.1), cd)
gf.name = "Class distribution chart"
ch = gf.chart
style_chart(ch)
ch.has_title = True
ch.chart_title.text_frame.text = "Images per class"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
ch.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONT
ch.has_legend = False
ch.category_axis.reverse_order = True
ch.category_axis.format.line.fill.background()
ch.category_axis.has_major_gridlines = False
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 9500
plot = ch.plots[0]
plot.gap_width = 40
plot.has_data_labels = True
plot.data_labels.number_format = "#,##0"
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(17)
plot.data_labels.font.name = FONT
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = rgb(DARK)
shp, tf = card(s, X0, 5.95, WIDTH, 0.95, "Imbalance note", fill=RED_TINT)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
paragraphs(tf, [[("67% of the images are common moles, ", {"bold": True, "color": RED}),
                 "so accuracy alone is misleading."]], 19, after=0)

# ================================================================== slide 9: implementation
s = new_slide("Implementation",
              "This is the web app. The user uploads an image. The app shows the top predictions, a separate melanoma "
              "probability, a heatmap and plain-language guidance. It warns when confidence is low and says that a "
              "non-melanoma result is not a clearance. Images are processed in memory and never stored.")
from PIL import Image  # noqa: E402

crop = Image.open(FIG / "fig3_3_app_result.png").convert("RGB")
crop.crop((0, 100, crop.width, 1240)).save(FIG / "ppt_app_result_top.png")
picture(s, "ppt_app_result_top.png", X0 + 0.4, 1.65, 6.6, "App result screen",
        "Screenshot of the web application result: prediction, melanoma meter, Grad-CAM heatmap")
_, tf = textbox(s, X0, 6.1, WIDTH, 1.0, "App features")
paragraphs(tf, [
    ["Top 3 classes, melanoma meter, heatmap and guidance"],
    ["Warns when unsure; images are never stored"],
], 19, bullet=True, after=5)

# ================================================================== slide 10: results, model comparison
s = new_slide("Results: Model Comparison",
              "Transfer learning made the biggest difference. Macro-F1 rose from 0.38 for the simple CNN and 0.40 for "
              "ResNet-18 trained from scratch to 0.67 with pretraining. Among the optimizers, Adam, RMSProp and SGD all "
              "worked; SGD with momentum failed at the learning rate we used. Augmentation helped most among the "
              "regularization methods, and class weights improve recall on the rare classes.")
cd = CategoryChartData()
cd.categories = ["CNN (scratch)", "ResNet-18 (scratch)", "ResNet-18 (pretrained)"]
cd.add_series("Macro-F1", (0.377, 0.402, 0.672))
cd.add_series("Balanced accuracy", (0.475, 0.505, 0.722))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X0), Inches(1.55), Inches(WIDTH), Inches(3.7), cd)
gf.name = "Model comparison chart"
ch = gf.chart
style_chart(ch)
ch.has_title = False
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(17)
ch.legend.font.name = FONT
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.minimum_scale = 0
ch.value_axis.maximum_scale = 0.9
ch.category_axis.format.line.color.rgb = rgb(GREY)
plot = ch.plots[0]
plot.gap_width = 60
plot.overlap = -5
plot.has_data_labels = True
plot.data_labels.number_format = "0.00"
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(17)
plot.data_labels.font.name = FONT
for ser, col in zip(plot.series, (DARK, ACCENT)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(col)
_, tf = textbox(s, X0, 5.4, WIDTH, 1.7, "Findings")
paragraphs(tf, [
    [("Pretraining: ", {"bold": True}), "macro-F1 from 0.40 to 0.67"],
    [("Optimizers: ", {"bold": True}), "Adam, RMSProp, SGD worked; momentum failed"],
    [("Also: ", {"bold": True}), "augmentation helped most; class weights help rare classes"],
], 18, bullet=True, after=5)

# ================================================================== slide 11: results, final model
s = new_slide("Results: Final Model",
              "Our selected model is ResNet-18 trained with RMSProp, chosen on the validation set. On the test set it "
              "reaches 79.2 percent accuracy, 0.69 balanced accuracy and 0.68 macro-F1. The heatmaps focus on the lesion. "
              "But it finds only 59 percent of melanomas; 21 percent are called moles. So the app never treats a "
              "non-melanoma result as a clearance.")
stats = [("79.2%", "accuracy", DARK), ("69%", "balanced accuracy", DARK), ("0.68", "macro-F1", DARK), ("59%", "melanomas found", RED)]
for i, (big, small, col) in enumerate(stats):
    x = X0 + (i % 2) * 1.95
    y = 1.65 + (i // 2) * 1.75
    shp, tf = card(s, x, y, 1.8, 1.6, f"Result tile {i + 1}", fill=RED_TINT if col == RED else TINT)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraphs(tf, [[big]], 30, after=0, align=PP_ALIGN.CENTER, bold=True, color=col)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    write(p, small, 17)
picture(s, "fig4_5_confusion_matrix.png", 4.4, 1.65, 3.5, "Confusion matrix", "Normalised confusion matrix of the final model on the test set")
_, tf = textbox(s, 4.4, 4.65, 3.5, 0.4, "Matrix caption")
paragraphs(tf, [["Confusion matrix (test set)"]], 17, after=0, color=GREY)
_, tf = textbox(s, X0, 5.3, WIDTH, 1.8, "Final model findings")
paragraphs(tf, [
    ["Heatmaps focus on the lesion itself"],
    [("Weak point: ", {"bold": True, "color": RED}), "21% of melanomas are predicted as moles"],
    [("A non-melanoma result is not a clearance", {"bold": True})],
], 18, bullet=True, after=6)

# ================================================================== slide 12: conclusion
s = new_slide("Conclusion",
              "To conclude: a lesion-wise split and balanced metrics gave honest results, and transfer learning mattered "
              "most. The final model reaches 79 percent accuracy but misses about four in ten melanomas, so it is an "
              "educational tool and not a medical device. For future work we would run several seeds, tune the "
              "optimizers, improve melanoma recall, calibrate the probabilities and test on other datasets.")
blocks = [
    ("Summary of outcomes", ["Transfer learning: macro-F1 0.40 to 0.67", "Final model: 79.2% accuracy, 59% of melanomas found", "Web app with heatmap, guidance and warnings"]),
    ("Objectives achieved", ["All five objectives achieved"]),
    ("Scope of improvement", ["Several seeds and tuned optimizers", "Better melanoma recall and calibrated probabilities", "Test on other datasets"]),
]
_, tf = textbox(s, X0, 1.6, WIDTH, 5.5, "Conclusion points")
first = True
for head, pts in blocks:
    par = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    par.space_before = Pt(8)
    par.space_after = Pt(3)
    write(par, head, 22, bold=True, color=DARK)
    for pt in pts:
        par = tf.add_paragraph()
        par.space_after = Pt(3)
        write(par, pt, 19)
        set_bullet(par)

# ================================================================== slide 13: references
s = new_slide("References",
              "These are the references cited in the slides. The full list is in the report.")
refs = [
    "[1] A. Esteva et al., “Dermatologist-level classification of skin cancer with deep neural networks,” Nature, vol. 542, pp. 115–118, 2017, doi: 10.1038/nature21056.",
    "[2] P. Tschandl, C. Rosendahl and H. Kittler, “The HAM10000 dataset, a large collection of multi-source dermatoscopic images of common pigmented skin lesions,” Sci. Data, vol. 5, 180161, 2018, doi: 10.1038/sdata.2018.161.",
    "[3] N. Codella et al., “Skin lesion analysis toward melanoma detection 2018: A challenge hosted by the International Skin Imaging Collaboration (ISIC),” arXiv:1902.03368, 2019.",
    "[4] K. He et al., “Deep residual learning for image recognition,” in Proc. IEEE CVPR, 2016, pp. 770–778, doi: 10.1109/CVPR.2016.90.",
    "[5] R. R. Selvaraju et al., “Grad-CAM: Visual explanations from deep networks via gradient-based localization,” in Proc. IEEE ICCV, 2017, pp. 618–626 (CVF Open Access).",
    "[6] J. Deng et al., “ImageNet: A large-scale hierarchical image database,” in Proc. IEEE CVPR, 2009, pp. 248–255, doi: 10.1109/CVPR.2009.5206848.",
]
_, tf = textbox(s, X0, 1.55, WIDTH, 5.6, "Reference list")
paragraphs(tf, [[r] for r in refs], 17, after=7)

# ================================================================== thank-you slide (template slide 14)
old_slides[13].notes_slide.notes_text_frame.text = "Thank you. We are happy to take questions."

# ------------------------------------------------------------------ order the slides, drop the unused template slides
sld_ids = list(prs.slides._sldIdLst)
by_part = {prs.part.related_part(i.rId): i for i in sld_ids}
new_ids = [i for i in sld_ids if prs.part.related_part(i.rId) not in {s_.part for s_ in old_slides}]
order = [by_part[title_slide.part]] + new_ids + [by_part[old_slides[13].part]]
for i in sld_ids:
    prs.slides._sldIdLst.remove(i)
for i in order:
    prs.slides._sldIdLst.append(i)
for s_ in old_slides:
    if s_ in (title_slide, old_slides[13]):
        continue
    for i in list(prs.slides._sldIdLst):
        if prs.part.related_part(i.rId) is s_.part:
            prs.part.drop_rel(i.rId)
            prs.slides._sldIdLst.remove(i)

prs.core_properties.title = "Skin Disease Detection using Deep Learning"
prs.core_properties.subject = "Deep Learning Lab (CSL701) mini project presentation"
prs.save(str(OUT))
print("saved", OUT, "slides:", len(prs.slides))
